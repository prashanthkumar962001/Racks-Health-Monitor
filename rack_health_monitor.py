"""
RackSense Daily Monitoring Script v4
--------------------------------------
- Opens Chrome browser
- You type username and password manually
- After login, script automatically scrapes all racks
- Pings each IP address
- Saves to dated Excel report

FIRST TIME SETUP (run once in PowerShell):
  pip install selenium webdriver-manager openpyxl beautifulsoup4

HOW TO RUN:
  python racksense_monitor.py
  Chrome will open -> type your login -> script does the rest
"""

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import subprocess
import platform
import time
import re
import sys

# ─────────────────────────────────────────
BASE_URL  = "http://rack_health_monitor-URL"
LOGIN_URL = f"{BASE_URL}/login"
RACKS_URL = f"{BASE_URL}/racks"
# ─────────────────────────────────────────


def ping_ip(ip):
    try:
        if platform.system().lower() == "windows":
            cmd = ["ping", "-n", "2", "-w", "1000", ip]
        else:
            cmd = ["ping", "-c", "2", "-W", "1", ip]
        result = subprocess.run(cmd, stdout=subprocess.DEVNULL,
                                stderr=subprocess.DEVNULL, timeout=6)
        return "Reachable" if result.returncode == 0 else "Unreachable"
    except Exception:
        return "Unreachable"


def start_browser():
    print("Starting Chrome browser...")
    options = Options()
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--ignore-certificate-errors")
    options.add_argument("--log-level=3")
    options.add_experimental_option("excludeSwitches", ["enable-logging"])
    driver = webdriver.Chrome(
        service=Service(ChromeDriverManager().install()),
        options=options
    )
    driver.maximize_window()
    return driver


def wait_for_login(driver):
    """
    Open login page, wait for user to type credentials and login.
    Script will automatically continue once /racks page is detected.
    """
    print("\n" + "=" * 56)
    print("  Chrome is open.")
    print("  Please type your Username and Password in the browser.")
    print("  Script will continue automatically after you log in.")
    print("=" * 56)

    driver.get(LOGIN_URL)

    # Wait until URL changes to /racks (meaning login succeeded)
    print("\nWaiting for you to log in...")
    timeout = 120  # give 2 minutes to login
    start = time.time()

    while time.time() - start < timeout:
        current_url = driver.current_url
        if "racks" in current_url or (
            "login" not in current_url and BASE_URL in current_url
        ):
            print("Login detected. Continuing...")
            return True
        time.sleep(1)

    print("\nERROR: Timed out waiting for login (2 minutes).")
    print("Please run the script again and login faster.")
    driver.quit()
    sys.exit(1)


def scrape_all_racks(driver):
    """Go to /racks and scrape every page."""
    print("\nNavigating to Racks page...")
    driver.get(RACKS_URL)
    time.sleep(3)

    all_racks = []
    page = 1

    while True:
        print(f"  Reading page {page}...")
        time.sleep(2)

        racks = scrape_current_page(driver)
        print(f"    Found {len(racks)} racks on page {page}")
        all_racks.extend(racks)

        # Try clicking Next button
        try:
            next_btn = driver.find_element(
                By.XPATH,
                "//a[normalize-space(text())='Next' and not(contains(@class,'disabled'))]"
                " | //button[normalize-space(text())='Next' and not(contains(@class,'disabled'))]"
                " | //li[not(contains(@class,'disabled'))]/a[normalize-space(text())='Next']"
            )
            if next_btn.is_enabled() and next_btn.is_displayed():
                driver.execute_script("arguments[0].click();", next_btn)
                page += 1
                time.sleep(2)
            else:
                break
        except Exception:
            break

    return all_racks


def scrape_current_page(driver):
    """Scrape rack rows visible on current page."""
    racks = []

    try:
        WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.XPATH, "//*[contains(text(),'AVAILABILITY')]"))
        )
    except Exception:
        pass

    time.sleep(1)
    soup = BeautifulSoup(driver.page_source, "html.parser")

    # Try standard selectors first
    rows = (
        soup.select("div.rack-item") or
        soup.select("div[class*='rack-item']") or
        soup.select("div[class*='rack-row']") or
        soup.select("div[class*='rack-card']") or
        soup.select("li[class*='rack']") or
        soup.select("tbody tr")
    )

    # Fallback: find any div containing IP + AVAILABILITY
    if not rows:
        for d in soup.find_all("div"):
            txt = d.get_text(" ", strip=True)
            if (re.search(r'\bIP\s+\d{1,3}\.\d{1,3}', txt, re.I) and
                    "AVAILABILITY" in txt.upper() and
                    len(d.find_all("div")) < 20):
                rows.append(d)

    seen = set()
    for row in rows:
        rack = extract_rack(row)
        if rack["rack_name"] and rack["rack_name"] not in seen:
            seen.add(rack["rack_name"])
            racks.append(rack)

    return racks


def extract_rack(el):
    text = el.get_text(" ", strip=True)
    rack = {"rack_name": "", "rack_id": "", "ip_address": "",
            "status": "", "availability": "", "location": "", "ping": ""}

    # Rack Name
    for tag in ["h5", "h4", "h3", "h2", "strong", "b"]:
        found = el.find(tag)
        if found:
            name = found.get_text(strip=True)
            if len(name) > 5:
                rack["rack_name"] = name
                break
    if not rack["rack_name"]:
        for sel in ["[class*='name']", "[class*='title']", "a"]:
            found = el.select_one(sel)
            if found:
                name = found.get_text(strip=True)
                if len(name) > 5:
                    rack["rack_name"] = name
                    break

    # ID
    m = re.search(r'\bID\s+(\d+)', text, re.I)
    if m:
        rack["rack_id"] = m.group(1)

    # IP Address
    m = re.search(r'\bIP\s+(\d{1,3}(?:\.\d{1,3}){3})\b', text, re.I)
    if not m:
        m = re.search(r'\b(\d{1,3}(?:\.\d{1,3}){3})\b', text)
    if m:
        rack["ip_address"] = m.group(1)

    # Availability
    m = re.search(r'AVAILABILITY\s*([\d.]+\s*%)', text, re.I)
    if m:
        rack["availability"] = m.group(1).replace(" ", "")

    # Status
    badge = (el.select_one(".badge") or
             el.select_one("[class*='badge']") or
             el.select_one("[class*='status']") or
             el.select_one("[class*='chip']"))
    if badge:
        rack["status"] = badge.get_text(strip=True)
    else:
        for word in ["Online", "Offline", "Unknown"]:
            if re.search(rf'\b{word}\b', text, re.I):
                rack["status"] = word
                break

    # Location
    loc = (el.select_one("small") or
           el.select_one("[class*='location']") or
           el.select_one("[class*='subtitle']") or
           el.select_one("p"))
    if loc:
        loc_text = loc.get_text(strip=True)
        if "/" in loc_text:
            rack["location"] = loc_text

    return rack


def export_to_excel(racks, timestamp):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rack Report"

    NAVY="1F3864"; WHITE="FFFFFF"; ALT="EBF2FF"
    GREEN="C6EFCE"; GREEN_FG="276221"
    RED="FFC7CE";   RED_FG="9C0006"
    YELLOW="FFEB9C"; ORANGE="FCE4D6"; LIGHT="DEEAF1"

    thin   = Side(style="thin", color="BDD7EE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value     = f"RackSense Daily Monitoring Report  -  {timestamp.strftime('%d-%m-%Y   %H:%M')}"
    c.font      = Font(bold=True, size=14, color=WHITE, name="Arial")
    c.fill      = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    # Summary
    total       = len(racks)
    online      = sum(1 for r in racks if r["status"].lower() == "online")
    unknown     = sum(1 for r in racks if "unknown" in r["status"].lower())
    offline     = sum(1 for r in racks if r["status"].lower() == "offline")
    reachable   = sum(1 for r in racks if r["ping"] == "Reachable")
    unreachable = sum(1 for r in racks if r["ping"] == "Unreachable")

    ws.merge_cells("A2:H2")
    c = ws["A2"]
    c.value = (f"Total: {total}   |   Online: {online}   |   Unknown: {unknown}   |   "
               f"Offline: {offline}   |   Reachable: {reachable}   |   Unreachable: {unreachable}")
    c.font      = Font(bold=True, size=10, name="Arial", color=NAVY)
    c.fill      = PatternFill("solid", fgColor=LIGHT)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # Headers
    cols   = ["#", "Rack Name", "Location", "ID", "IP Address",
              "Status", "Availability", "Ping Status"]
    widths = [5,   50,          38,          8,    18, 12, 14, 15]

    for ci, (h, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(row=3, column=ci, value=h)
        c.font      = Font(bold=True, color=WHITE, name="Arial", size=10)
        c.fill      = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = border
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[3].height = 20

    # Data rows
    for i, rack in enumerate(racks, 1):
        row = i + 3
        bg  = ALT if i % 2 == 0 else WHITE
        vals = [i, rack["rack_name"], rack["location"], rack["rack_id"],
                rack["ip_address"], rack["status"], rack["availability"], rack["ping"]]

        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=ci, value=val)
            c.font      = Font(name="Arial", size=9)
            c.alignment = Alignment(vertical="center",
                                    horizontal="left" if ci == 2 else "center")
            c.border    = border

            if ci == 6:
                sl = str(val).lower()
                if sl == "online":
                    c.fill = PatternFill("solid", fgColor=GREEN)
                    c.font = Font(name="Arial", size=9, color=GREEN_FG, bold=True)
                elif "unknown" in sl:
                    c.fill = PatternFill("solid", fgColor=YELLOW)
                    c.font = Font(name="Arial", size=9, color="7D6608", bold=True)
                elif sl == "offline":
                    c.fill = PatternFill("solid", fgColor=RED)
                    c.font = Font(name="Arial", size=9, color=RED_FG, bold=True)
                else:
                    c.fill = PatternFill("solid", fgColor=bg)
            elif ci == 7:
                try:
                    pct = float(str(val).replace("%", ""))
                    if pct >= 95:   c.fill = PatternFill("solid", fgColor=GREEN)
                    elif pct >= 50: c.fill = PatternFill("solid", fgColor=YELLOW)
                    elif pct == 0:  c.fill = PatternFill("solid", fgColor=RED)
                    else:           c.fill = PatternFill("solid", fgColor=ORANGE)
                except Exception:
                    c.fill = PatternFill("solid", fgColor=bg)
            elif ci == 8:
                if val == "Reachable":
                    c.fill = PatternFill("solid", fgColor=GREEN)
                    c.font = Font(name="Arial", size=9, color=GREEN_FG, bold=True)
                elif val == "Unreachable":
                    c.fill = PatternFill("solid", fgColor=RED)
                    c.font = Font(name="Arial", size=9, color=RED_FG, bold=True)
                else:
                    c.fill = PatternFill("solid", fgColor=bg)
            else:
                c.fill = PatternFill("solid", fgColor=bg)

        ws.row_dimensions[row].height = 16

    ws.freeze_panes = "A4"
    fname = f"Rack_Report_{timestamp.strftime('%d-%m-%Y')}.xlsx"
    wb.save(fname)
    return fname


def main():
    print("=" * 56)
    print("        RackSense Daily Monitor")
    print("=" * 56)
    print(f"  Site : {BASE_URL}")

    driver = start_browser()

    try:
        wait_for_login(driver)

        print("\nScraping rack data...")
        racks = scrape_all_racks(driver)

    finally:
        driver.quit()
        print("Browser closed.")

    if not racks:
        print("\nWARNING: No racks were found.")
        print("Please share a screenshot with your developer.")
        sys.exit(1)

    print(f"\nFound {len(racks)} racks. Pinging IPs now...\n")

    for i, rack in enumerate(racks, 1):
        ip = rack.get("ip_address", "")
        if ip:
            rack["ping"] = ping_ip(ip)
            result = ">>" if rack["ping"] == "Reachable" else "XX"
            print(f"  [{i:>3}/{len(racks)}]  {ip:<18}  {result}  {rack['ping']}")
        else:
            rack["ping"] = "No IP"

    timestamp = datetime.now()
    fname = export_to_excel(racks, timestamp)

    print(f"\n{'=' * 56}")
    print(f"  Report saved   : {fname}")
    print(f"  Total racks    : {len(racks)}")
    print(f"  Online         : {sum(1 for r in racks if r['status'].lower() == 'online')}")
    print(f"  Reachable IPs  : {sum(1 for r in racks if r.get('ping') == 'Reachable')}")
    print(f"  Date / Time    : {timestamp.strftime('%d-%m-%Y  %H:%M:%S')}")
    print(f"{'=' * 56}")
    input("\n  Press Enter to close...")


if __name__ == "__main__":
    main()
  
